from __future__ import annotations

import argparse
import hashlib
import re
import json
import uuid
from datetime import datetime, timezone
from pathlib import Path

import duckdb
import pandas as pd
import yaml
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
BRONZE_DIR = ROOT / "bronze_files"
SILVER_DIR = ROOT / "silver"
CONFIG_PATH = ROOT / "config" / "silver_config.yaml"
REGISTRY_DB = SILVER_DIR / "silver_registry.duckdb"
LOGBOOK_PATH = SILVER_DIR / "silver_logbook.md"
HUD_NOTICE_REF = "Notice H 2025-01"

# ── final output column sets (spec-only) ──────────────────────────────────────

QCT_FINAL_COLS = [
    # Identifiers
    "designation_id", "source_pk",
    # Geography
    "tract_fips", "county_fips", "state_fips", "cbsa_code", "fmr_area_code",
    "is_metro_tract", "area_population", "split_tr_flag",
    # Designation
    "is_designated", "designation_year", "hud_notice_ref",
    "basis_boost_pct", "qct_trigger_criterion", "effective_date",
    # Economic criteria — current vintage
    "poverty_rate_at_designation", "income_criterion_ratio",
    "median_hh_income", "median_hh_income_moe", "VLIL4_current",
    # Prior vintages (2018+; NULL pre-2018 — expected gap)
    "poverty_rate_vintage_minus1", "poverty_rate_vintage_minus2",
    "income_criterion_ratio_minus1", "income_criterion_ratio_minus2",
    # Same-row trend signals (2018+; NULL pre-2018 — expected gap)
    "poverty_rate_trend_3yr", "poverty_trend_direction", "income_ratio_trend_3yr",
    # ACS reliability (2018+; NULL pre-2018 — expected gap)
    "income_estimate_cv", "income_estimate_reliable",
]

# ZCTA-level detail table (feeds SDDADesignation aggregation)
SDDA_ZCTA_COLS = [
    "zcta_code", "fmr_area_code", "area_name", "is_designated", "ranking_ratio",
    "safmr_2br", "vlil_4person", "total_population", "population_in_qct",
    "effective_population", "designation_year", "lihtc_max_rent",
    "qct_overlap_pct", "is_territory", "effective_date",
]

# Area-level aggregation — the SDDADesignation node
SDDA_AREA_COLS = [
    "designation_id", "fmr_area_code", "area_name", "is_designated",
    "ranking_ratio", "safmr_2br", "vlil_4person", "total_population",
    "population_in_qct", "effective_population", "zcta_count_designated",
    "designation_year", "basis_boost_pct", "effective_date",
    "lihtc_max_rent", "qct_overlap_pct",
]

# NMDDADesignation node
NMDDA_FINAL_COLS = [
    "designation_id", "county_fips", "fmr_area_code", "area_name",
    "is_designated", "ranking_ratio", "fmr_2br", "vlil_4person",
    "total_population", "population_in_qct", "effective_population",
    "designation_year", "basis_boost_pct", "is_territory",
    "effective_date", "lihtc_max_rent", "qct_overlap_pct",
]

_LOGBOOK_HEADER = (
    "# Silver Logbook\n\n"
    "| Timestamp (UTC) | Dataset | Year | Status | File | Rows | Reason |\n"
    "|---|---|---|---|---|---|---|\n"
)


def _ts() -> str:
    return datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")


def logbook_row(
    dataset: str,
    year: int,
    status: str,
    filename: str = "",
    rows: int = 0,
    reason: str = "",
) -> None:
    SILVER_DIR.mkdir(parents=True, exist_ok=True)
    if not LOGBOOK_PATH.exists():
        LOGBOOK_PATH.write_text(_LOGBOOK_HEADER, encoding="utf-8")
    row_str = f"| {_ts()} | {dataset} | {year} | {status} | {filename} | {rows:,} | {reason} |\n"
    with LOGBOOK_PATH.open("a", encoding="utf-8") as f:
        f.write(row_str)


CONFIG_KEYS = {
    "qct_static_keep_columns",
    "qct_column_map",
    "qct_required",
    "qct_fingerprint_fields",
    "dda_stable_column_map",
    "dda_metro_fingerprint_fields",
    "dda_nonmetro_fingerprint_fields",
}


def load_config(path: Path = CONFIG_PATH) -> dict:
    if not path.exists():
        raise FileNotFoundError(path)
    with path.open("r", encoding="utf-8") as handle:
        config = yaml.safe_load(handle) or {}
    missing = sorted(CONFIG_KEYS - set(config))
    if missing:
        raise ValueError(f"Silver config missing keys: {', '.join(missing)}")
    return config


CONFIG = load_config()


# ── helpers ───────────────────────────────────────────────────────────────────

def clean_code(value: object, width: int) -> str | None:
    if pd.isna(value):
        return None
    text = str(value).strip()
    if text.endswith(".0"):
        text = text[:-2]
    if not text or text.lower() in {"nan", "none"}:
        return None
    digits = "".join(ch for ch in text if ch.isdigit())
    return digits.zfill(width) if digits else None


def _clean_cbsa(value: object) -> str | None:
    """Normalise cbsa_code: strip .0 float artifact, remove sentinels, zfill(5)."""
    if pd.isna(value):
        return None
    s = str(value).strip()
    if s.endswith(".0"):
        s = s[:-2]
    if not s or s.lower() in {"nan", "none", "100000"}:
        return None
    return s.zfill(5) if s.isdigit() else None


def _enforce_types(df: pd.DataFrame) -> pd.DataFrame:
    """
    Enforce canonical dtypes before writing any Silver parquet.

    Boolean cols  → native Python bool (pandas bool_ dtype)
    Integer cols  → Int64 (nullable)
    Float cols    → float64
    String FIPS   → zero-padded str
    Null policy   → None/NaN, never empty string

    Only acts on columns that exist in df.
    """
    _BOOL_COLS = {
        "is_designated", "is_territory", "is_high_disparity",
        "is_multifamily_constrained", "income_estimate_reliable",
        "split_tr_flag", "is_metro_tract",
    }
    _INT_COLS = {
        "designation_year", "basis_boost_pct", "assessment_year",
        "year", "ami_pct", "mfi_value", "state_mfi",
        "max_income", "max_rent",
        "limit_1person", "limit_2person", "limit_3person", "limit_4person",
        "limit_5person", "limit_6person", "limit_7person", "limit_8person",
        "total_applications", "total_originated", "total_denied",
        "multifamily_applications", "multifamily_denied", "lender_count",
        "area_population", "total_population", "population_in_qct",
        "effective_population",
    }
    _FLOAT_COLS = {
        "poverty_rate_at_designation", "poverty_rate_vintage_minus1",
        "poverty_rate_vintage_minus2", "income_criterion_ratio",
        "income_criterion_ratio_minus1", "income_criterion_ratio_minus2",
        "poverty_rate_trend_3yr", "income_ratio_trend_3yr",
        "income_estimate_cv", "ranking_ratio", "qct_overlap_pct",
        "denial_rate_overall", "denial_rate_minority", "denial_rate_white",
        "denial_rate_disparity_ratio", "multifamily_denial_rate",
        "fair_lending_risk_score", "gse_sold_pct", "avg_loan_amount",
        "safmr_2br", "vlil_4person", "fmr_2br", "lihtc_max_rent",
        "median_hh_income", "median_hh_income_moe",
    }
    _STR_NULL_SENTINELS = {"", "NA", "nan", "NaN", "null", "None", "NULL"}
    _STR_COLS = {
        "area_name", "county_name", "cbsa_title", "metro_name",
        "qct_trigger_criterion", "poverty_trend_direction",
        "top_denial_reason", "risk_tier", "program_type", "hud_notice_ref",
        "hud_fmr_area_code", "rent_metric_type",
    }

    for col in list(df.columns):
        if col in _BOOL_COLS:
            df[col] = (
                pd.to_numeric(df[col], errors="coerce").fillna(0).astype(bool)
            )
        elif col in _INT_COLS:
            df[col] = pd.to_numeric(df[col], errors="coerce").round().astype("Int64")
        elif col in _FLOAT_COLS:
            df[col] = pd.to_numeric(df[col], errors="coerce")
        elif col in _STR_COLS:
            df[col] = df[col].apply(
                lambda x: None if (pd.isna(x) or str(x).strip() in _STR_NULL_SENTINELS)
                else str(x).strip()
            )
        elif col == "cbsa_code":
            df[col] = df[col].apply(_clean_cbsa)
        elif col in {"tract_fips", "fips_code"}:
            df[col] = df[col].apply(lambda x: clean_code(x, 11))
        elif col == "county_fips":
            df[col] = df[col].apply(lambda x: clean_code(x, 5))
        elif col == "state_fips":
            df[col] = df[col].apply(lambda x: clean_code(x, 2))

    return df


def to_numeric(series: pd.Series) -> pd.Series:
    cleaned = series.astype(str).str.replace(r"[$,]", "", regex=True)
    return pd.to_numeric(cleaned, errors="coerce")


def to_int(series: pd.Series) -> pd.Series:
    return to_numeric(series).round().astype("Int64")


def source_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def fingerprint_row(row: pd.Series, fields: list[str]) -> str:
    data = {field: "" if pd.isna(row[field]) else str(row[field]) for field in fields}
    canonical = json.dumps(data, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(canonical.encode("utf-8")).hexdigest()


def write_registry(
    conn: duckdb.DuckDBPyConnection,
    dataset: str,
    year: int,
    path: str,
    df: pd.DataFrame,
    bronze_source_id: str,
) -> None:
    conn.execute(
        """
        create table if not exists snapshot_registry (
            dataset          varchar,
            version          varchar,
            path             varchar,
            row_count        integer,
            schema_hash      varchar,
            content_hash     varchar,
            bronze_source_id varchar,
            written_at       varchar
        )
        """
    )
    conn.execute(
        "delete from snapshot_registry where dataset = ? and version = ? and path = ?",
        [dataset, str(year), path],
    )
    schema_hash = hashlib.sha256(json.dumps(sorted(df.columns.tolist())).encode("utf-8")).hexdigest()
    content = "".join(df["content_hash"].fillna("").astype(str).tolist()) if "content_hash" in df else ""
    row = {
        "dataset": dataset,
        "version": str(year),
        "path": path,
        "row_count": len(df),
        "schema_hash": schema_hash,
        "content_hash": hashlib.sha256(content.encode("utf-8")).hexdigest(),
        "bronze_source_id": bronze_source_id,
        "written_at": pd.Timestamp.now("UTC").isoformat(),
    }
    row_df = pd.DataFrame([row])
    conn.register("_reg_row", row_df)
    conn.execute("insert into snapshot_registry select * from _reg_row")
    conn.unregister("_reg_row")


def _write_dedup(df: pd.DataFrame, pk_col: str, out_path: Path) -> None:
    """Write parquet, merging with any existing file and deduplicating on pk_col.
    New rows take priority -- existing rows with the same pk are overwritten.
    """
    if out_path.exists():
        existing = pd.read_parquet(out_path)
        existing = existing[~existing[pk_col].isin(df[pk_col])]
        combined = pd.concat([existing, df], ignore_index=True)
    else:
        combined = df.copy()
    combined = combined.drop_duplicates(subset=[pk_col], keep="last").reset_index(drop=True)
    combined.to_parquet(out_path, index=False)


def require_columns(df: pd.DataFrame, required: set[str], label: str) -> None:
    missing = sorted(required - set(df.columns))
    if missing:
        raise ValueError(f"{label} missing required columns: {', '.join(missing)}")


def _sheet_names(path: Path) -> list[str]:
    """Return sheet names, falling back to calamine if openpyxl rejects the file."""
    if path.suffix.lower() != ".xlsx":
        return pd.ExcelFile(path, engine="xlrd").sheet_names
    try:
        return load_workbook(path, read_only=True).sheetnames
    except TypeError:
        return pd.ExcelFile(path, engine="calamine").sheet_names


def load_sheet(path: Path, sheet_name: str) -> pd.DataFrame:
    """Read one sheet, falling back to calamine if openpyxl rejects the file."""
    if path.suffix.lower() != ".xlsx":
        engine = "xlrd"
    else:
        try:
            df = pd.read_excel(path, sheet_name=sheet_name, engine="openpyxl")
            df.columns = df.columns.astype(str).str.strip()
            if df.empty:
                raise ValueError(f"{path.name} sheet {sheet_name!r} loaded zero rows")
            return df
        except TypeError:
            engine = "calamine"
    df = pd.read_excel(path, sheet_name=sheet_name, engine=engine)
    if df.empty:
        raise ValueError(f"{path.name} sheet {sheet_name!r} loaded zero rows")
    df.columns = df.columns.astype(str).str.strip()
    return df


# ── column resolvers ──────────────────────────────────────────────────────────

def _find(cols: list[str], candidates: list[str], label: str) -> str:
    """Case-insensitive column lookup; returns the original-case column name."""
    lower_map = {c.lower(): c for c in cols}
    for c in candidates:
        orig = lower_map.get(c.lower())
        if orig is not None:
            return orig
    raise ValueError(f"None of {candidates!r} found in {label} columns")


def _find_opt(cols: list[str], candidates: list[str]) -> str | None:
    """Case-insensitive optional lookup; returns None if no candidate found."""
    lower_map = {c.lower(): c for c in cols}
    for c in candidates:
        orig = lower_map.get(c.lower())
        if orig is not None:
            return orig
    return None


def _apply_column_map(raw: pd.DataFrame, full_map: dict) -> pd.DataFrame:
    """Keep only columns that appear in full_map (case-insensitive), then rename."""
    lower_map = {c.lower(): c for c in raw.columns}
    resolved: dict[str, str] = {}
    for src, dst in full_map.items():
        orig = lower_map.get(src.lower())
        if orig is not None and orig not in resolved:
            resolved[orig] = dst
    keep = list(resolved.keys())
    return raw[keep].rename(columns=resolved)


_QCT_AUX_SHEETS = {"variable definitions", "state codes", "county codes", "metro cbsa codes"}

# Multi-year bundle filenames — mirrors ingest_QCTDDA._QCT_BUNDLE
_QCT_BUNDLE: dict[int, str] = {}
for _y in range(2003, 2007): _QCT_BUNDLE[_y] = "qct_data_2003_2004_2005_2006.xlsx"
for _y in range(2007, 2010): _QCT_BUNDLE[_y] = "qct_data_2007_2008_2009.xlsx"
for _y in range(2010, 2013): _QCT_BUNDLE[_y] = "qct_data_2010_2011_2012.xlsx"
for _y in range(2013, 2015): _QCT_BUNDLE[_y] = "qct_data_2013_2014.xls"

# DDA xlsx availability: 2016 and below only have PDFs -- routed to silver_dda_pdf.py
_DDA_MIN_YEAR = 2017


def _qct_source(year: int) -> Path:
    """Resolve the bronze QCT file for a given year, handling multi-year bundles."""
    year_dir = BRONZE_DIR / "qct" / str(year)
    individual = year_dir / f"qct_data_{year}.xlsx"
    if individual.exists():
        return individual
    bundle_name = _QCT_BUNDLE.get(year)
    if bundle_name:
        candidate = year_dir / bundle_name
        if candidate.exists():
            return candidate
    raise FileNotFoundError(
        f"No QCT bronze file for {year}. Expected: {individual}"
        + (f" or {year_dir / bundle_name}" if bundle_name else "")
    )


def _qct_data_sheets(sheetnames: list[str]) -> list[str]:
    """
    Returns the 1 or 2 data sheet names to load.
    Most years have an AL sheet + MT sheet; 2023 ships a single combined sheet.
    """
    al = next((s for s in sheetnames if s.upper().startswith("AL")), None)
    mt = next((s for s in sheetnames if s.upper().startswith("MT")), None)
    if al and mt:
        return [al, mt]
    data = [s for s in sheetnames if s.lower().strip() not in _QCT_AUX_SHEETS]
    if len(data) == 1:
        return data
    raise ValueError(f"Cannot detect QCT data sheets from: {sheetnames}")


def resolve_qct_columns(df_columns: list, designation_year: int) -> dict:
    """
    Maps actual source column names → canonical internal names.

    ACS vintage: designation_year - 3  (2025→22, 2024→21, ... 2015→12)
    2015 files use ICPercent_NN / PovertyRate_NN instead of inc_factor / pov_rate.
    2016-2017 files have no B19013 median-income columns (ACS suppressed).
    All lookups are case-insensitive; original-case column name is returned.
    """
    cols = [str(c).strip() for c in df_columns]
    v0 = str(designation_year - 3)[2:]   # e.g. "22" for 2025
    v1 = str(designation_year - 4)[2:]
    v2 = str(designation_year - 5)[2:]

    def find(candidates):
        return _find(cols, candidates, "QCT")

    col_map = {
        # tract FIPS: 2015+ use tract_id_NN; 2013-2014 use tract_id; 2003-2012 use fips
        find(["tract_id", f"tract_id_{v0}", f"tract_id_{v1}", "tract_id_20", "fips"]): "tract_fips",
        # FMR area: cbsasub with year suffix (2018+), plain (2007-2017); absent 2003-2006
        _find_opt(cols, [f"cbsasub{v0}", "cbsasub24", "cbsasub"]) or "__missing_cbsasub__": "fmr_area_code",
        # Poverty rate: versioned (2016+), no-suffix (2013-2014), legacy 'povrate' (2003-2012)
        find([f"pov_rate_{v0}", f"PovertyRate_{v0}", "pov_rate", "povrate"]): "poverty_rate_at_designation",
        # Income criterion ratio: versioned (2016+), no-suffix (2013-2014), legacy 'pctinc' (2003-2012)
        find([f"inc_factor_{v0}", f"ICPercent_{v0}", "inc_factor", "pctinc"]): "income_criterion_ratio",
        # VLIL4: versioned (2018+), no-suffix (2013-2014), legacy 'amgi60pc' (2007-2012)
        find([
            f"VLIL4_{designation_year - 3}", f"VLIL4_{v0}",
            f"VLIL_4_Person_{v0}", f"VLIL_4_Person_{designation_year - 3}",
            "VLIL_4_Person", "amgi60pc", "pr4_60pc", "trinclim",
        ]): "VLIL4_current",
        # P1 total pop: Census 2020 PL, Census 2010 SF1, legacy Census 2000
        find(["P1", "p0010001", "pop100"]): "P1",
        # P5 group quarters: Census 2020 PL, Census 2010 SF1, Census 2000 P87, 2003-2006 direct
        find(["P5", "p0420001", "p087002", "p015001", "groupqtr"]): "P5",
    }
    # H1 housing units: Census 2020 PL, Census 2010 SF1, Census 2000 — absent in 2003-2006
    h1_src = _find_opt(cols, ["H1", "p0150001", "p016001"])
    if h1_src:
        col_map[h1_src] = "H1"
    # Remove spurious __missing__ key (cbsasub absent in 2003-2006)
    col_map.pop("__missing_cbsasub__", None)

    # Vintage-minus1/minus2 trend columns — optional; absent in pre-2015 files
    pov_v1 = _find_opt(cols, [f"pov_rate_{v1}", f"PovertyRate_{v1}"])
    pov_v2 = _find_opt(cols, [f"pov_rate_{v2}", f"PovertyRate_{v2}"])
    icr_v1 = _find_opt(cols, [f"inc_factor_{v1}", f"ICPercent_{v1}"])
    icr_v2 = _find_opt(cols, [f"inc_factor_{v2}", f"ICPercent_{v2}"])
    if pov_v1: col_map[pov_v1] = "poverty_rate_vintage_minus1"
    if pov_v2: col_map[pov_v2] = "poverty_rate_vintage_minus2"
    if icr_v1: col_map[icr_v1] = "income_criterion_ratio_minus1"
    if icr_v2: col_map[icr_v2] = "income_criterion_ratio_minus2"

    # Median income: versioned ACS (2018+), no-suffix (2013-2014), legacy medfaminc (2003-2006)
    b19013_src = _find_opt(cols, [f"B19013est1_{v0}", "B19013est1", "medfminc"])
    b19013me_src = _find_opt(cols, [f"B19013me1_{v0}", "B19013me1"])
    if b19013_src:
        col_map[b19013_src] = "median_hh_income"
    if b19013me_src:
        col_map[b19013me_src] = "median_hh_income_moe"

    return col_map


def resolve_dda_metro_columns(df_columns: list, year: int) -> dict:
    """
    All years use ZCTA/SAFMR format for the metro sheet.
    Year-varying: QCT ref year, SAFMR vintage, VLIL vintage, designation flag column.
    """
    cols = [str(c).strip() for c in df_columns]
    fy = year - 1  # SAFMR/VLIL vintage year

    def find(candidates):
        return _find(cols, candidates, "DDA metro")

    pop_gte = next((c for c in cols if "Pop>=" in c or "Pop >" in c), None)
    pyear = year - 1  # prior QCT year (some vintages reference year-1 QCT)
    col_map: dict[str, str] = {
        find(["ZIP Code Tabulation Area (ZCTA)", "ZCTA"]): "zcta_code",
        # 2017 used "FY{fy} Adjusted"; 2018+ use "{fy} Final 40th Percentile"
        find([
            f"{fy} Final 40th Percentile 2-Bedroom SAFMR",
            f"FY{fy} Adjusted 2-Bedroom SAFMR",
            f"FY{fy} Final 40th Percentile 2-Bedroom SAFMR",
        ]): "safmr_2br",
        # 2017 references prior QCT year; 2018+ reference current year
        find([
            f"ZCTA Population also in {year} QCT",
            f"Population also in {year} QCT",
            f"ZCTA Population also in {pyear} QCT",
            f"Population also in {pyear} QCT",
        ]): "population_in_qct",
        find([
            f"ZCTA Population NOT in {year} QCT",
            f"Population NOT in {year} QCT",
            f"ZCTA Population NOT in {pyear} QCT",
            f"Population NOT in {pyear} QCT",
        ]): "effective_population",
        # 2017 uses "FY{fy}" prefix; 2018+ omit "FY"
        find([
            f"{fy} 4-Person VLIL",
            f"{fy} 4-Person Very Low Income Limit (VLIL)",
            f"FY{fy} 4-Person VLIL",
        ]): "vlil_4person",
        find([f"{year} SDDA (1=SDDA)"]): "is_designated",
    }
    if pop_gte:
        col_map[pop_gte] = "pop_gte_100_flag"
    return col_map


def resolve_dda_nonmetro_columns(df_columns: list, year: int) -> dict:
    """
    Nonmetro DDA columns. FIPS is optional (absent 2018-2020).
    'Effective Pop' is absent pre-2022; fall back to 'Population NOT in {year} QCT'.
    """
    cols = [str(c).strip() for c in df_columns]
    fy = year - 1

    def find(candidates):
        return _find(cols, candidates, "DDA nonmetro")

    col_map: dict[str, str] = {
        find([f"FY{fy} Final 40th Percentile 2-Bedroom FMR"]): "fmr_2br",
        find([f"Population also in {year} QCT", f"Population also in {year-1} QCT", "Pop in QCT"]): "population_in_qct",
        find(["Effective Pop", f"Population NOT in {year} QCT", f"Population NOT in {year-1} QCT"]): "effective_population",
        find([f"{fy} 4-Person Very Low Income Limit (VLIL)", f"{fy} 4-Person VLIL"]): "vlil_4person",
        find([f"{year} NMDDA"]): "is_designated",
    }

    # FIPS absent in 2018-2020; optional
    fips_src = _find_opt(cols, ["FIPS"])
    if fips_src:
        col_map[fips_src] = "county_fips"

    return col_map


# ── QCT logic ─────────────────────────────────────────────────────────────────

def compute_qct_criterion(row: pd.Series) -> str:
    income_met = pd.notna(row["income_criterion_ratio"]) and row["income_criterion_ratio"] < 1.0
    poverty_met = pd.notna(row["poverty_rate_at_designation"]) and row["poverty_rate_at_designation"] >= 0.25
    if row["is_designated"] == 1:
        if income_met and poverty_met:
            return "both"
        if poverty_met:
            return "poverty"
        if income_met:
            return "income"
        return "area_constraint"
    return "none"


def poverty_direction(trend: object) -> str:
    if pd.isna(trend):
        return "unknown"
    if trend < -0.05:
        return "improving"
    if trend > 0.05:
        return "worsening"
    return "stable"


def run_qct(year: int) -> None:
    qct_dir = SILVER_DIR / "qct"
    qct_dir.mkdir(parents=True, exist_ok=True)
    effective_date = f"{year}-01-01"
    source = _qct_source(year)  # handles individual files and multi-year bundles
    is_bundle = source.name in _QCT_BUNDLE.values()

    data_sheets = _qct_data_sheets(_sheet_names(source))

    raw = pd.concat(
        [load_sheet(source, s) for s in data_sheets],
        ignore_index=True,
    )

    # Bundle files contain multiple years — filter to the requested year.
    # The year column is typically named "Year" or embedded in a column like "Designation_Year".
    if is_bundle:
        year_col = _find_opt(raw.columns.tolist(), ["Year", "Designation_Year", "designation_year"])
        if year_col:
            raw = raw[raw[year_col].astype(str).str.strip() == str(year)].copy()
            print(f"QCT {year}: filtered bundle to {len(raw):,} rows for year {year}")
        else:
            print(f"QCT {year}: WARNING — bundle file has no year column; all {len(raw):,} rows loaded")

    # Row-count floor: lower threshold for older/bundle-filtered years
    min_rows = 5000 if is_bundle else 40000
    if len(raw) <= min_rows:
        raise ValueError(f"QCT structural row count too low: {len(raw)} (min {min_rows})")
    # Resolve designation column: prefer plain 'qct', fall back to year-specific 'qct_{year}'
    cols_lower_map = {c.lower(): c for c in raw.columns}
    if "qct" not in cols_lower_map:
        year_qct = f"qct_{year}"
        if year_qct in cols_lower_map:
            # e.g. 'qct_2010' in 2010-2012 bundle — add to dynamic map
            extra_desig = {cols_lower_map[year_qct]: "is_designated"}
            print(f"QCT {year}: using designation column '{cols_lower_map[year_qct]}'")
        else:
            qct_year_cols = [c for c in cols_lower_map if re.fullmatch(r"qct_\d{4}", c)]
            if qct_year_cols:
                # 2011/2012 from 2010-2012 bundle only has 'qct_2010'; reuse it
                fallback = cols_lower_map[qct_year_cols[0]]
                extra_desig = {fallback: "is_designated"}
                print(f"QCT {year}: no '{year_qct}' found — using bundle column '{fallback}' as designation proxy")
            else:
                raise ValueError(f"QCT designation column 'qct' or '{year_qct}' not found in {list(raw.columns[:10])}...")
    else:
        extra_desig = {}

    dynamic_map = resolve_qct_columns(raw.columns.tolist(), year)
    full_map = {**CONFIG["qct_column_map"], **dynamic_map, **extra_desig}
    # _apply_column_map is case-insensitive; handles Area_pop vs Area_Pop etc.
    df = _apply_column_map(raw, full_map)

    require_columns(df, set(CONFIG["qct_required"]), "QCT")
    # VLIL4/P1/P5/H1 required for derived metrics; warn if absent in very old vintages
    for col in ["VLIL4_current", "P1", "P5", "H1"]:
        if col not in df.columns:
            if year >= 2015:
                raise ValueError(f"QCT missing derivation helper column '{col}' after resolve+rename")
            print(f"QCT {year}: '{col}' not found in old-format file — derived metrics will be null")

    df["tract_fips"] = df["tract_fips"].map(lambda x: clean_code(x, 11))
    df["county_fips"] = df["county_fips"].map(lambda x: clean_code(x, 5))
    df["state_fips"] = df["state_fips"].map(lambda x: clean_code(x, 2))
    df["source_pk"] = df["source_pk"].map(lambda x: clean_code(x, 12))
    # Keep as int during validation; cast to bool at write time
    df["is_designated"] = to_int(df["is_designated"])
    df["is_metro_tract"] = to_int(df["is_metro_tract"])

    # Drop blank Excel trailer rows where is_designated is NaN
    nan_desig = df["is_designated"].isna().sum()
    if nan_desig > 0:
        print(f"QCT {year}: dropping {nan_desig} blank trailer rows (null is_designated)")
        df = df[df["is_designated"].notna()].copy()

    # Ensure median income columns exist (absent in 2016-2017)
    for col in ["median_hh_income", "median_hh_income_moe"]:
        if col not in df.columns:
            df[col] = pd.NA

    numeric_cols = [
        "area_population",
        "poverty_rate_at_designation",
        "poverty_rate_vintage_minus1",
        "poverty_rate_vintage_minus2",
        "median_hh_income",
        "median_hh_income_moe",
        "income_criterion_ratio",
        "income_criterion_ratio_minus1",
        "income_criterion_ratio_minus2",
        "VLIL4_current",
        "P1",
        "P5",
    ]
    for col in numeric_cols:
        if col in df.columns:
            df[col] = to_numeric(df[col])
    df["area_population"] = df["area_population"].astype("Int64")

    null_icr = df["income_criterion_ratio"].isna().sum()
    if null_icr > 0:
        null_desig = df[df["income_criterion_ratio"].isna() & (df["is_designated"] == 1)]
        poverty_covered   = (null_desig["poverty_rate_at_designation"] >= 0.25).sum()
        area_constraint   = len(null_desig) - poverty_covered
        print(
            f"QCT NOTE: {null_icr:,} rows have null income_criterion_ratio "
            f"({len(null_desig):,} designated — {poverty_covered} poverty-criterion, "
            f"{area_constraint} area_constraint/ACS-suppressed)"
        )
    if not df["is_designated"].dropna().isin([0, 1]).all() or df["is_designated"].isna().any():
        raise ValueError("QCT is_designated must be binary with no nulls")
    if not df["tract_fips"].dropna().str.len().eq(11).all():
        raise ValueError("QCT tract_fips values must be 11 characters")

    row_ok = (
        df["tract_fips"].fillna("").str.fullmatch(r"\d{11}")
        & df["county_fips"].fillna("").str.fullmatch(r"\d{5}")
        & df["is_designated"].isin([0, 1])
    )
    df["_row_status"] = row_ok.map({True: "KEEP", False: "SKIP"})
    bad_tract   = ~df["tract_fips"].fillna("").str.fullmatch(r"\d{11}")
    bad_county  = ~df["county_fips"].fillna("").str.fullmatch(r"\d{5}")
    bad_desig   = ~df["is_designated"].isin([0, 1])
    df["_skip_reason"] = (
        bad_tract.map({True: "bad_tract_fips;",  False: ""})
        + bad_county.map({True: "bad_county_fips;", False: ""})
        + bad_desig.map({True: "bad_is_designated;", False: ""})
    )

    df["designation_id"] = df["tract_fips"].astype(str) + f"_{year}"
    df["basis_boost_pct"] = df["is_designated"].map(lambda x: 30 if x == 1 else 0)
    # Fix 1: normalise cbsa_code — strip .0 float artifact, remove 100000 sentinel
    if "cbsa_code" in df.columns:
        df["cbsa_code"] = df["cbsa_code"].apply(_clean_cbsa)
    df["designation_year"] = year
    df["qct_trigger_criterion"] = df.apply(compute_qct_criterion, axis=1)
    # Trend/derived columns — null when vintage-minus columns absent (pre-2015 files)
    pov_v2 = df.get("poverty_rate_vintage_minus2")
    icr_v2 = df.get("income_criterion_ratio_minus2")
    df["poverty_rate_trend_3yr"] = (df["poverty_rate_at_designation"] - pov_v2) if pov_v2 is not None else pd.NA
    df["poverty_trend_direction"] = df["poverty_rate_trend_3yr"].map(poverty_direction) if pov_v2 is not None else "unknown"
    df["income_ratio_trend_3yr"] = (df["income_criterion_ratio"] - icr_v2) if icr_v2 is not None else pd.NA
    moe = df.get("median_hh_income_moe")
    if moe is not None:
        # ACS MOE is at 90% CI; SE = MOE / 1.645; CV = SE / estimate
        df["income_estimate_cv"] = (moe.abs() / 1.645) / df["median_hh_income"].replace(0, pd.NA)
        df["income_estimate_reliable"] = (df["income_estimate_cv"] < 0.30).fillna(False)
    else:
        df["income_estimate_cv"] = pd.NA
        df["income_estimate_reliable"] = pd.NA
    # ci_spans_income_threshold requires AMI cutoff from a separate dataset — computed in Gold
    p1 = df.get("P1")
    p5 = df.get("P5")
    if p1 is not None and p5 is not None:
        df["group_quarters_pct"] = (p5 / p1.replace(0, pd.NA)).fillna(0)
    else:
        df["group_quarters_pct"] = 0.0
    df["is_gq_dominated"] = df["group_quarters_pct"] > 0.50
    df["effective_date"] = effective_date
    df["expiry_date"] = None
    df["hud_notice_ref"] = HUD_NOTICE_REF
    df["bronze_source_id"] = source_hash(source)
    df["pipeline_run_id"] = str(uuid.uuid4())

    if (df.loc[df["is_designated"] == 1, "basis_boost_pct"] != 30).any():
        raise ValueError("QCT basis_boost_pct invariant failed for designated rows")
    if (df.loc[df["is_designated"] == 0, "basis_boost_pct"] != 0).any():
        raise ValueError("QCT basis_boost_pct invariant failed for non-designated rows")
    if "split_tr_flag" not in df.columns:
        raise ValueError("split_tr_flag column missing after rename")

    split_rows = df[df["split_tr_flag"] == 1][["tract_fips", "county_fips", "state_fips", "source_pk", "split_tr_flag"]]
    split_rows.to_csv(SILVER_DIR / "silver_qa_qct_split_tracts.csv", index=False)
    bad_split = split_rows[~split_rows["state_fips"].isin(["23", "25"])]
    if not bad_split.empty:
        bad_split.to_csv(SILVER_DIR / "silver_qa_qct_split_tracts_unexpected_state.csv", index=False)

    df.loc[df["is_gq_dominated"], ["tract_fips", "county_fips", "group_quarters_pct"]].to_csv(
        SILVER_DIR / "silver_qa_gq_dominated.csv", index=False
    )
    df = df.drop(columns=["group_quarters_pct", "is_gq_dominated", "P1", "P5", "H1"], errors="ignore")

    # Ensure all fingerprint fields exist before hashing (median_hh_income may be absent 2016-2017)
    for field in CONFIG["qct_fingerprint_fields"]:
        if field not in df.columns:
            df[field] = pd.NA
    df["content_hash"] = df.apply(lambda row: fingerprint_row(row, CONFIG["qct_fingerprint_fields"]), axis=1)

    pq_path = qct_dir / f"silver_qct_{year}.parquet"
    skipped_path = qct_dir / f"silver_qct_{year}_skipped.parquet"
    # Skipped parquet retains all diagnostic cols for QA review
    df[df["_row_status"] == "SKIP"].to_parquet(skipped_path, index=False)
    # Main parquet: spec columns only, with enforced dtypes
    out_cols = [c for c in QCT_FINAL_COLS if c in df.columns]
    out_df = df[out_cols].copy()
    # designation_id uniqueness check — warn and deduplicate rather than abort
    dupes = out_df[out_df.duplicated("designation_id", keep=False)]
    if not dupes.empty:
        print(f"  WARN QCT {year}: {len(dupes)} duplicate designation_ids — keeping last")
        out_df = out_df.drop_duplicates("designation_id", keep="last").reset_index(drop=True)
    # Fix 2: enforce canonical types before write
    out_df = _enforce_types(out_df)
    out_df.to_parquet(pq_path, index=False)

    with duckdb.connect(str(REGISTRY_DB)) as conn:
        write_registry(conn, "qct", year, str(pq_path.relative_to(SILVER_DIR)), df, df["bronze_source_id"].iloc[0])

    print(f"Wrote QCT Silver: {pq_path} ({len(df):,} rows)")
    logbook_row("QCT", year, "✅ OK", pq_path.name, len(df))


# ── DDA logic ─────────────────────────────────────────────────────────────────

def mark_dda_status(df: pd.DataFrame, code_col: str) -> pd.DataFrame:
    code_width = 5
    row_ok = (
        df[code_col].fillna("").str.fullmatch(r"\d{" + str(code_width) + r"}")
        & df["is_designated"].isin([0, 1])
        & df["ranking_ratio"].notna()
    )
    df["_row_status"] = row_ok.map({True: "KEEP", False: "SKIP"})
    df["_skip_reason"] = ""
    df.loc[~df[code_col].fillna("").str.fullmatch(r"\d{" + str(code_width) + r"}"), "_skip_reason"] += f"bad_{code_col};"
    df.loc[~df["is_designated"].isin([0, 1]), "_skip_reason"] += "bad_is_designated;"
    df.loc[df["ranking_ratio"].isna(), "_skip_reason"] += "missing_ranking_ratio;"
    return df


def _build_nmdda_county_fips_lookup() -> dict[str, str]:
    """
    Fix 6: Build (state_abbr, county_name_lower) → county_fips lookup from
    Census ANSI county file (national_county.txt).  Used to fill missing
    county_fips on NMDDA rows where HUD didn't include them (2017-2020).
    """
    ref = ROOT / "bronze_files" / "Geographic" / "national_county.txt"
    if not ref.exists():
        return {}
    try:
        nt = pd.read_csv(ref, sep="|", dtype=str)
        lookup: dict[str, str] = {}
        for _, row in nt.iterrows():
            state = row["STATE"].strip().upper()
            fips5 = row["STATEFP"].zfill(2) + row["COUNTYFP"].zfill(3)
            name_raw = row["COUNTYNAME"].strip()
            # Store full name and stripped suffixes
            lookup[(state, name_raw.lower())] = fips5
            for suf in (" County", " Parish", " Borough", " Census Area",
                        " Municipio", " Municipality", " city", " City"):
                if name_raw.endswith(suf):
                    lookup[(state, name_raw[:-len(suf)].lower())] = fips5
        return lookup
    except Exception:
        return {}


_NMDDA_FIPS_LOOKUP: dict[str, str] | None = None


def _fill_nmdda_county_fips(df: pd.DataFrame) -> pd.DataFrame:
    """
    Fix 6: Attempt to fill null county_fips from area_name via the county
    name lookup.  Only fills rows where county_fips is currently null.
    """
    global _NMDDA_FIPS_LOOKUP
    if _NMDDA_FIPS_LOOKUP is None:
        _NMDDA_FIPS_LOOKUP = _build_nmdda_county_fips_lookup()
    if not _NMDDA_FIPS_LOOKUP or "area_name" not in df.columns:
        return df

    _ABBR_TO_FIPS = {
        "AL": "01", "AK": "02", "AZ": "04", "AR": "05", "CA": "06",
        "CO": "08", "CT": "09", "DE": "10", "DC": "11", "FL": "12",
        "GA": "13", "HI": "15", "ID": "16", "IL": "17", "IN": "18",
        "IA": "19", "KS": "20", "KY": "21", "LA": "22", "ME": "23",
        "MD": "24", "MA": "25", "MI": "26", "MN": "27", "MS": "28",
        "MO": "29", "MT": "30", "NE": "31", "NV": "32", "NH": "33",
        "NJ": "34", "NM": "35", "NY": "36", "NC": "37", "ND": "38",
        "OH": "39", "OK": "40", "OR": "41", "PA": "42", "RI": "44",
        "SC": "45", "SD": "46", "TN": "47", "TX": "48", "UT": "49",
        "VT": "50", "VA": "51", "WA": "53", "WV": "54", "WI": "55",
        "WY": "56", "PR": "72", "VI": "78", "GU": "66", "AS": "60",
        "MP": "69",
    }

    def _lookup_fips(row: pd.Series) -> str | None:
        if pd.notna(row.get("county_fips")):
            return row["county_fips"]
        state = str(row.get("state_abbr", "") or "").strip().upper()
        name  = str(row.get("area_name",  "") or "").strip()
        if not state or not name:
            return None
        # Try full name, then strip common suffixes
        for candidate in (
            name,
            re.sub(r",.*$", "", name).strip(),  # strip ", StateAbbr" suffix
            re.sub(r"\s+(County|Parish|Borough|Census Area|Municipio).*$",
                   "", name, flags=re.I).strip(),
        ):
            hit = _NMDDA_FIPS_LOOKUP.get((state, candidate.lower()))
            if hit:
                return hit
        return None

    null_mask = df["county_fips"].isna()
    if null_mask.any():
        filled = df[null_mask].apply(_lookup_fips, axis=1)
        filled_count = filled.notna().sum()
        df.loc[null_mask, "county_fips"] = filled
        if filled_count:
            print(f"  Fix 6: filled {filled_count} null county_fips from area_name lookup")
    return df


def run_dda(year: int) -> None:
    if year < _DDA_MIN_YEAR:
        # Route pre-2017 years to the PDF parser
        from silver_dda_pdf import run_dda_from_pdf
        run_dda_from_pdf(year)
        return

    dda_dir = SILVER_DIR / "dda"
    dda_dir.mkdir(parents=True, exist_ok=True)
    effective_date = f"{year}-01-01"
    pipeline_run_id = str(uuid.uuid4())
    source = BRONZE_DIR / "dda" / str(year) / f"{year}-DDAs-Data-Used-to-Designate.xlsx"
    if not source.exists():
        raise FileNotFoundError(source)

    sheetnames = _sheet_names(source)
    # 2018-2021: "NonMetro" / "Metro"; 2022+: "NM DDA {year}" / "NMDDA" / "Metro DDA {year}"
    nonmetro_sheet = next(
        s for s in sheetnames
        if ("NM" in s and "DDA" in s) or "nonmetro" in s.lower()
    )
    metro_sheet = next(
        s for s in sheetnames
        if s != nonmetro_sheet and ("MDDA" in s or "metro" in s.lower())
    )

    metro_raw    = load_sheet(source, metro_sheet)
    nonmetro_raw = load_sheet(source, nonmetro_sheet)

    stable_map = CONFIG["dda_stable_column_map"]
    metro    = _apply_column_map(metro_raw,    {**stable_map, **resolve_dda_metro_columns(metro_raw.columns.tolist(), year)})
    nonmetro = _apply_column_map(nonmetro_raw, {**stable_map, **resolve_dda_nonmetro_columns(nonmetro_raw.columns.tolist(), year)})

    # fmr_area_code absent in 2018 metro (no CBSAsub column in that vintage)
    # county_fips absent in 2018-2020 nonmetro (no FIPS column in those vintages)
    require_columns(metro,    {"zcta_code", "is_designated", "ranking_ratio", "safmr_2br", "vlil_4person"}, "DDA metro")
    require_columns(nonmetro, {"is_designated", "ranking_ratio", "fmr_2br", "vlil_4person"}, "DDA nonmetro")

    territory_state_fips = {"60", "66", "69", "72", "78"}

    metro["area_type"] = "metro"
    nonmetro["area_type"] = "nonmetro"

    for df in [metro, nonmetro]:
        for col in ["total_population", "population_in_qct", "effective_population", "cumulative_pop"]:
            if col in df.columns:
                df[col] = to_int(df[col])
        for col in ["vlil_4person", "lihtc_max_rent_hud", "ranking_ratio", "cumulative_pct"]:
            if col in df.columns:
                df[col] = to_numeric(df[col])
        df["is_designated"] = to_int(df["is_designated"])
        df["lihtc_max_rent"] = df["vlil_4person"] * 0.03
        df["qct_overlap_pct"] = df["population_in_qct"] / df["total_population"].replace(0, pd.NA)
        df["designation_year"] = year
        df["basis_boost_pct"] = df["is_designated"].map(lambda x: 30 if x == 1 else 0)
        df["hud_notice_ref"] = HUD_NOTICE_REF
        df["effective_date"] = effective_date
        df["expiry_date"] = None
        df["bronze_source_id"] = source_hash(source)
        df["pipeline_run_id"] = pipeline_run_id

    # Drop blank Excel trailer rows with null is_designated
    nan_m = metro["is_designated"].isna().sum()
    if nan_m:
        print(f"DDA {year} metro: dropping {nan_m} blank trailer rows (null is_designated)")
    metro = metro[metro["is_designated"].notna()].copy()
    nan_nm = nonmetro["is_designated"].isna().sum()
    if nan_nm:
        print(f"DDA {year} nonmetro: dropping {nan_nm} blank trailer rows (null is_designated)")
    nonmetro = nonmetro[nonmetro["is_designated"].notna()].copy()

    metro["zcta_code"] = metro["zcta_code"].map(lambda x: clean_code(x, 5))
    metro["safmr_2br"] = to_numeric(metro["safmr_2br"])
    if "pop_gte_100_flag" in metro.columns:
        metro["pop_gte_100_flag"] = to_int(metro["pop_gte_100_flag"])
    _territory_zcta_prefixes = {"006", "007", "008", "009", "969"}  # PR, USVI, GU
    _territory_kws = {"puerto rico", "virgin islands", "guam", "american samoa",
                      "northern mariana", "palau", "micronesia", "marshall"}
    if "zcta_code" in metro.columns:
        metro["is_territory"] = metro["zcta_code"].astype(str).str[:3].isin(_territory_zcta_prefixes)
    elif "area_name" in metro.columns:
        metro["is_territory"] = metro["area_name"].astype(str).str.lower().apply(
            lambda n: any(t in n for t in _territory_kws)
        )
    else:
        metro["is_territory"] = False
    metro["rent_metric_type"] = "SAFMR"
    metro = mark_dda_status(metro, "zcta_code")

    if "county_fips" in nonmetro.columns:
        nonmetro["county_fips"] = nonmetro["county_fips"].map(lambda x: clean_code(x, 5))
        nonmetro["is_territory"] = nonmetro["county_fips"].map(
            lambda x: bool(str(x).isdigit() and str(x)[:2] in territory_state_fips)
        )
        nonmetro = mark_dda_status(nonmetro, "county_fips")
    else:
        # county_fips absent (2018-2020 vintages): detect territories by area_name keyword
        _territory_names = {"puerto rico", "virgin islands", "guam", "american samoa",
                            "northern mariana", "palau", "micronesia", "marshall"}
        if "area_name" in nonmetro.columns:
            nonmetro["is_territory"] = nonmetro["area_name"].astype(str).str.lower().apply(
                lambda n: any(t in n for t in _territory_names)
            )
        else:
            nonmetro["is_territory"] = False
        row_ok = nonmetro["is_designated"].isin([0, 1]) & nonmetro["ranking_ratio"].notna()
        nonmetro["_row_status"] = row_ok.map({True: "KEEP", False: "SKIP"})
        bad_desig = ~nonmetro["is_designated"].isin([0, 1])
        nonmetro["_skip_reason"] = bad_desig.map({True: "bad_is_designated;", False: ""})
        nonmetro.loc[nonmetro["ranking_ratio"].isna(), "_skip_reason"] += "missing_ranking_ratio;"

    if "fmr_2br" in nonmetro.columns:
        nonmetro["fmr_2br"] = to_numeric(nonmetro["fmr_2br"])
    nonmetro["rent_metric_type"] = "FMR"
    # NMDDADesignation: designation_id keyed on county_fips per spec
    if "county_fips" in nonmetro.columns:
        nonmetro["designation_id"] = nonmetro["county_fips"].astype(str) + f"_{year}"
    elif "area_name" in nonmetro.columns:
        nonmetro["designation_id"] = nonmetro["area_name"].astype(str).str.replace(" ", "_") + f"_{year}"
    else:
        nonmetro["designation_id"] = pd.NA

    if not metro["is_designated"].dropna().isin([0, 1]).all() or metro["is_designated"].isna().any():
        raise ValueError("DDA metro is_designated must be binary with no nulls")
    if not nonmetro["is_designated"].dropna().isin([0, 1]).all() or nonmetro["is_designated"].isna().any():
        raise ValueError("DDA nonmetro is_designated must be binary with no nulls")
    # HUD rounds LIHTC to the nearest dollar; allow up to $1.00 difference
    if "lihtc_max_rent_hud" in metro.columns:
        if not (metro["lihtc_max_rent"].sub(metro["lihtc_max_rent_hud"]).abs().dropna() < 1.0).all():
            raise ValueError("DDA metro LIHTC rent formula invariant failed")
    if "lihtc_max_rent_hud" in nonmetro.columns:
        if not (nonmetro["lihtc_max_rent"].sub(nonmetro["lihtc_max_rent_hud"]).abs().dropna() < 1.0).all():
            raise ValueError("DDA nonmetro LIHTC rent formula invariant failed")
    if nonmetro["is_territory"].any():
        undesig_territories = nonmetro.loc[nonmetro["is_territory"] & (nonmetro["is_designated"] != 1)]
        if not undesig_territories.empty:
            print(f"  WARN DDA {year}: {len(undesig_territories)} territory rows not designated — check source data")

    # Capture availability of optional geographic keys BEFORE padding with pd.NA
    metro_has_fmr_area = "fmr_area_code" in metro.columns
    nonmetro_has_fips  = "county_fips"   in nonmetro.columns

    # Pad missing fields before fingerprinting
    for field in CONFIG["dda_metro_fingerprint_fields"]:
        if field not in metro.columns:
            metro[field] = pd.NA
    for field in CONFIG["dda_nonmetro_fingerprint_fields"]:
        if field not in nonmetro.columns:
            nonmetro[field] = pd.NA

    metro["content_hash"]    = metro.apply(lambda row: fingerprint_row(row, CONFIG["dda_metro_fingerprint_fields"]), axis=1)
    nonmetro["content_hash"] = nonmetro.apply(lambda row: fingerprint_row(row, CONFIG["dda_nonmetro_fingerprint_fields"]), axis=1)

    # Fix 6: fill missing county_fips from area_name lookup (mainly 2017-2020)
    nonmetro = _fill_nmdda_county_fips(nonmetro)

    # Fix 2: cast bool fields and enforce canonical types before write
    metro    = _enforce_types(metro.copy())
    nonmetro = _enforce_types(nonmetro.copy())

    # designation_id uniqueness checks
    for label, df_check, pk in [("SDDA metro", metro, "zcta_code"),
                                  ("NMDDA", nonmetro, "designation_id")]:
        if pk in df_check.columns:
            dupes = df_check[df_check.duplicated(pk, keep=False)]
            if not dupes.empty:
                print(f"  WARN DDA {year} {label}: {len(dupes)} duplicate {pk} — deduplicating")

    metro_path    = dda_dir / f"silver_dda_{year}_metro_zcta.parquet"
    nonmetro_path = dda_dir / f"silver_dda_{year}_nonmetro_county.parquet"
    _write_dedup(metro[    [c for c in SDDA_ZCTA_COLS   if c in metro.columns]    ], "zcta_code",      metro_path)
    _write_dedup(nonmetro[ [c for c in NMDDA_FINAL_COLS if c in nonmetro.columns] ], "designation_id", nonmetro_path)
    tables: list[tuple[Path, pd.DataFrame]] = [(metro_path, metro), (nonmetro_path, nonmetro)]

    if metro_has_fmr_area:
        dda1_metro = metro[metro["is_designated"] == 1].copy()
        area_agg = dda1_metro.groupby("fmr_area_code").agg(
            area_name=("area_name", "first"),
            is_designated=("is_designated", "max"),
            ranking_ratio=("ranking_ratio", "max"),
            safmr_2br=("safmr_2br", "max"),
            vlil_4person=("vlil_4person", "first"),
            total_population=("total_population", "sum"),
            population_in_qct=("population_in_qct", "sum"),
            effective_population=("effective_population", "sum"),
            zcta_count_designated=("zcta_code", "count"),
        ).reset_index()
        area_agg["lihtc_max_rent"]   = area_agg["vlil_4person"] * 0.03
        area_agg["qct_overlap_pct"]  = area_agg["population_in_qct"] / area_agg["total_population"].replace(0, pd.NA)
        area_agg["rent_metric_type"] = "SAFMR"
        area_agg["area_type"]        = "metro"
        area_agg["is_territory"]     = False
        area_agg["designation_year"] = year
        area_agg["designation_id"]   = area_agg["fmr_area_code"].astype(str) + f"_{year}"
        area_agg["basis_boost_pct"]  = 30
        area_agg["hud_notice_ref"]   = HUD_NOTICE_REF
        area_agg["effective_date"]   = effective_date
        area_agg["expiry_date"]      = None
        area_agg["bronze_source_id"] = source_hash(source)
        area_agg["pipeline_run_id"]  = pipeline_run_id
        area_agg["_row_status"]      = "KEEP"
        area_agg["_skip_reason"]     = ""
        area_fields = [f for f in CONFIG["dda_nonmetro_fingerprint_fields"] if f != "county_fips"]
        for field in area_fields:
            if field not in area_agg.columns:
                area_agg[field] = pd.NA
        area_agg["content_hash"] = area_agg.apply(lambda row: fingerprint_row(row, area_fields), axis=1)
        area_path = dda_dir / f"silver_dda_{year}_metro_area.parquet"
        _write_dedup(area_agg[[c for c in SDDA_AREA_COLS if c in area_agg.columns]], "designation_id", area_path)
        tables.append((area_path, area_agg))
        area_count = len(area_agg)
    else:
        print(f"DDA {year}: skipping metro area_agg (no fmr_area_code in metro ZCTA data)")
        area_count = 0

    with duckdb.connect(str(REGISTRY_DB)) as conn:
        for pq_path, df in tables:
            write_registry(conn, "dda", year, str(pq_path.relative_to(SILVER_DIR)), df, df["bronze_source_id"].iloc[0])

    total_rows = len(metro) + len(nonmetro) + area_count
    print(
        f"Wrote DDA Silver: {dda_dir} "
        f"({len(metro):,} metro ZCTA, {len(nonmetro):,} nonmetro county, {area_count:,} metro area rows)"
    )
    logbook_row("DDA", year, "✅ OK", dda_dir.name, total_rows)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Build Silver QCT/DDA Parquet snapshots from bronze_files.",
        epilog=(
            "Examples:\n"
            "  python silver_qct_dda.py --year 2025\n"
            "  python silver_qct_dda.py --year 2022 2023 2024 2025\n"
            "  python silver_qct_dda.py --year-range 2017 2025\n"
            "  python silver_qct_dda.py --year-range 2017 2025 --dataset qct\n"
        ),
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    parser.add_argument("--dataset", choices=["all", "qct", "dda"], default="all")
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument(
        "--year", type=int, nargs="+",
        help="One or more designation years (e.g. --year 2022 2023 2024)",
    )
    group.add_argument(
        "--year-range", type=int, nargs=2, metavar=("FROM", "TO"),
        help="Inclusive year range (e.g. --year-range 2017 2025)",
    )
    args = parser.parse_args()
    if args.year_range:
        lo, hi = args.year_range
        if lo > hi:
            parser.error(f"--year-range FROM ({lo}) must be <= TO ({hi})")
        args.year = list(range(lo, hi + 1))
    return args


def main() -> None:
    args = parse_args()
    for year in args.year:
        print(f"\n--- Year {year} ---")
        if args.dataset in {"all", "qct"}:
            try:
                run_qct(year)
            except FileNotFoundError as e:
                print(f"  QCT SKIP: {e}")
                logbook_row("QCT", year, "⚠️ SKIP", reason=str(e))
            except Exception as e:
                print(f"  QCT FAIL: {e}")
                logbook_row("QCT", year, "❌ FAIL", reason=str(e))
        if args.dataset in {"all", "dda"}:
            try:
                run_dda(year)
            except FileNotFoundError as e:
                print(f"  DDA SKIP: {e}")
                logbook_row("DDA", year, "⚠️ SKIP", reason=str(e))
            except Exception as e:
                print(f"  DDA FAIL: {e}")
                logbook_row("DDA", year, "❌ FAIL", reason=str(e))


if __name__ == "__main__":
    main()
